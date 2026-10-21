#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Самопроверка DoctorHide.

Собирает фикстуры (md, docx, pdf с текстом, скан без текста, png) и прогоняет их
через настоящий CLI, проверяя результат: что скрыто, что сохранено, читается ли файл
и не осталось ли данных под плашками.

Запуск:
    python tests/selftest.py
    python tests/selftest.py -v      # подробно
"""

from __future__ import annotations

import json
import os
import subprocess
import sys

import pymupdf
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP = os.path.join(ROOT, "doctorhide.py")
WORK = os.path.join(ROOT, "_private", "_selftest")

VERBOSE = "-v" in sys.argv

DOC_MD = """Договор оказания услуг № 12/2024

Исполнитель: ООО «Пример», ИНН 7707083893, КПП 770701001, ОГРН 1027700132195
Адрес: 125009, г. Москва, ул. Тверская, д. 15, кв. 42
Директор: Иванов Иван Иванович, тел. +7 (495) 123-45-67, e-mail: ivanov@example.ru
Сотрудник Петров Пётр Петрович, 35 лет, паспорт 4509 123456, СНИЛС 112-233-445 95
Банк: ПАО «Сбербанк», р/с 40702810123456789012, БИК 044525225
Выручка за квартал: 1 200 000 руб.
"""

COURT_MD = """РЕШЕНИЕ
Именем Российской Федерации

Кировский районный суд г. Москвы в составе председательствующего судьи
Сидорова А.В., при секретаре Кузнецовой М.И., рассмотрев дело № 2-123/2023
по иску Иванова И.И. к Петрову П.П., 45 лет, проживающего по адресу:
г. Москва, ул. Ленина, д. 5, кв. 10, с участием прокурора Смирнова А.А.
установил:
"""

SCHOOL_MD = """Директору ГБОУ СОШ № 42 Центрального района г. Твери Петровой А.И.
от Сергеева-Кузнецова Дмитрия Олеговича, паспорт 2812 445566

Заявление

Прошу разрешить Сергееву-Кузнецову Артёму Дмитриевичу, ученику 7 Г класса,
в связи с его участием в мероприятиях в Театре юношеского творчества Тверского
городского Дворца творчества юных ранний уход с занятий 15 апреля 2026 года.
"""

BANNED_DOC = ["Иванов", "Петров", "7707083893", "770701001", "1027700132195",
              "125009", "Москва", "Тверская", "123-45-67", "ivanov@", "4509",
              "123456", "112-233-445", "1 200 000", "40702810123456789012"]

results: list[tuple[bool, str, str]] = []


def check(ok: bool, name: str, detail: str = "") -> None:
    results.append((bool(ok), name, detail))
    mark = "OK  " if ok else "FAIL"
    print(f"[{mark}] {name}" + (f" — {detail}" if detail and (VERBOSE or not ok) else ""))


def run(args, stdin_text=None):
    return subprocess.run([sys.executable, APP] + args, cwd=WORK, capture_output=True,
                          text=True, encoding="utf-8", errors="replace", input=stdin_text)


def pdf_lines(path, lines, fontsize=11):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    font = r"C:\Windows\Fonts\arial.ttf"
    page.insert_font(fontname="F0", fontfile=font)
    y = 60
    for line in lines:
        if line:
            page.insert_text((60, y), line, fontname="F0", fontsize=fontsize)
        y += 18
    doc.save(path)
    return path


def pdf_text(path):
    doc = pymupdf.open(path)
    text = "".join(p.get_text() for p in doc)
    pages = doc.page_count
    meta = {k: v for k, v in doc.metadata.items() if v}
    doc.close()
    return text, pages, meta


def pdf_streams(path):
    """Видимый текст: сырые байты файла и потоки содержимого страниц.

    Шрифтовые программы и растры не проверяем — там встречаются те же цифры
    в служебных таблицах глифов, это не утечка.
    """
    doc = pymupdf.open(path)
    blobs = [open(path, "rb").read()]
    for page in doc:
        try:
            blobs.append(page.read_contents() or b"")
        except Exception:
            pass
    doc.close()
    return b"\n".join(blobs).decode("utf-8", "replace")


def whited_pixels(src, dst):
    a = Image.open(src).convert("L")
    b = Image.open(dst).convert("L")
    if a.size != b.size:
        return -1
    pa, pb = a.load(), b.load()
    w, h = a.size
    count = 0
    for y in range(0, h, 2):
        for x in range(0, w, 2):
            if pa[x, y] < 128 <= pb[x, y]:
                count += 1
    return count


def main() -> int:
    os.makedirs(WORK, exist_ok=True)
    for name in os.listdir(WORK):
        try:
            os.remove(os.path.join(WORK, name))
        except OSError:
            pass

    print("== окружение ==")
    check(os.path.exists(APP), "программа на месте", APP)
    tess = os.path.join(ROOT, "_tools", "tessdata", "rus.traineddata")
    check(os.path.exists(tess), "русская модель Tesseract найдена", tess)

    with open(os.path.join(WORK, "doc.md"), "w", encoding="utf-8") as f:
        f.write(DOC_MD)
    with open(os.path.join(WORK, "court.md"), "w", encoding="utf-8") as f:
        f.write(COURT_MD)

    with open(os.path.join(WORK, "school.md"), "w", encoding="utf-8") as f:
        f.write(SCHOOL_MD)

    from docx import Document

    d = Document()
    d.add_heading("Акт проверки", level=1)
    d.add_paragraph("Организация: АО «Вектор», ИНН 7701234567")
    d.add_paragraph("Директор: Смирнов Сергей Сергеевич, 52 года.")
    d.save(os.path.join(WORK, "doc.docx"))

    lines = DOC_MD.splitlines()
    pdf_lines(os.path.join(WORK, "doc.pdf"), lines)

    src = pymupdf.open(os.path.join(WORK, "doc.pdf"))
    pix = src[0].get_pixmap(dpi=200)
    png = os.path.join(WORK, "doc.png")
    pix.save(png)
    scan = pymupdf.open()
    sp = scan.new_page(width=595, height=842)
    sp.insert_image(sp.rect, filename=png)
    scan.save(os.path.join(WORK, "scan.pdf"))
    scan.close()
    src.close()

    print("\n== md ==")
    r = run(["doc.md"])
    out_md = os.path.join(WORK, "doc_anon.md")
    check(r.returncode == 0, "код возврата 0", r.stderr.strip()[:120])
    check(os.path.exists(out_md), "создан doc_anon.md")
    if os.path.exists(out_md):
        text = open(out_md, encoding="utf-8").read()
        left = [b for b in BANNED_DOC if b in text]
        check(not left, "персональные данные скрыты", f"осталось: {left}")
        check("[censored]" in text, "маркер [censored] на месте")
        check("Договор оказания услуг" in text, "неизменный текст сохранён")

    print("\n== pdf с текстовым слоем ==")
    r = run(["doc.pdf"])
    out_pdf = os.path.join(WORK, "doc_anon.pdf")
    check(r.returncode == 0, "код возврата 0", r.stderr.strip()[:120])
    check(os.path.exists(out_pdf), "создан doc_anon.pdf")
    if os.path.exists(out_pdf):
        text, pages, meta = pdf_text(out_pdf)
        left = [b for b in BANNED_DOC if b in text]
        check(not left, "в тексте нет скрытых данных", f"осталось: {left}")
        check(pages == 1, "число страниц сохранено", str(pages))
        check(len(text) > 100, "документ остаётся читаемым", f"{len(text)} символов")
        meta_core = {k: v for k, v in meta.items() if k != "format"}
        check(not meta_core, "метаданные очищены", str(meta_core))
        streams = pdf_streams(out_pdf)
        left_s = [b for b in BANNED_DOC if b in streams]
        check(not left_s, "в потоках PDF нет скрытых данных", f"осталось: {left_s}")

    print("\n== PDF всегда остаётся PDF ==")
    r = run(["doc.pdf", "-o", "wrong.md"])
    fixed = os.path.join(WORK, "wrong.pdf")
    check(os.path.exists(fixed), "результат сохранён как .pdf", f"код {r.returncode}")
    if os.path.exists(fixed):
        head = open(fixed, "rb").read(5)
        check(head == b"%PDF-", "внутри действительно PDF", repr(head))
        check(not os.path.exists(os.path.join(WORK, "wrong.md")), "файл с чужим расширением не создан")

    print("\n== судебный акт ==")
    r = run(["court.md"])
    out_court = os.path.join(WORK, "court_anon.md")
    check(r.returncode == 0, "код возврата 0", r.stderr.strip()[:120])
    if os.path.exists(out_court):
        text = open(out_court, encoding="utf-8").read()
        kept = ["Сидорова А.В.", "Кузнецовой М.И.", "Смирнова А.А.", "Иванова И.И.", "Кировский районный суд", "2-123/2023"]
        missed = [k for k in kept if k not in text]
        check(not missed, "суд и участники сохранены", f"пропало: {missed}")
        hidden = [x for x in ["Москва", "45 лет"] if x in text]
        check(not hidden, "адрес/город/возраст скрыты", f"осталось: {hidden}")

    print("\n== школьное заявление ==")
    r = run(["school.md"])
    out_school = os.path.join(WORK, "school_anon.md")
    check(r.returncode == 0, "код возврата 0", r.stderr.strip()[:120])
    if os.path.exists(out_school):
        text = open(out_school, encoding="utf-8").read()
        banned = ["Петровой", "Сергеев", "Кузнецов", "2812", "445566", "ГБОУ", "СОШ",
                  "Театр", "Твери"]
        left = [b for b in banned if b in text]
        check(not left, "школа, ФИО, паспорт и театр скрыты", f"осталось: {left}")
        check("Заявление" in text and "7 Г класса" in text, "текст заявления сохранён")

    print("\n== docx ==")
    r = run(["doc.docx"])
    out_docx = os.path.join(WORK, "doc_anon.md")
    check(r.returncode == 0, "код возврата 0", r.stderr.strip()[:120])
    if os.path.exists(out_docx):
        text = open(out_docx, encoding="utf-8").read()
        check("# Акт проверки" in text, "заголовок перенесён в Markdown")
        check("7701234567" not in text, "ИНН скрыт")
        check("Смирнов" not in text, "ФИО скрыто")

    print("\n== скан без текстового слоя ==")
    r = run(["scan.pdf"])
    out_scan = os.path.join(WORK, "scan_anon.pdf")
    check(r.returncode == 0, "код возврата 0", r.stderr.strip()[:120])
    if os.path.exists(out_scan):
        text, pages, _ = pdf_text(out_scan)
        check(pages == 1, "страница сохранена", str(pages))
        check(len(text.strip()) > 50, "текстовый слой добавлен", f"{len(text)} символов")
        left = [b for b in BANNED_DOC if b in text]
        check(not left, "в текстовом слое нет скрытых данных", f"осталось: {left}")
        shot = os.path.join(WORK, "scan_anon.png")
        pymupdf.open(out_scan)[0].get_pixmap(dpi=200).save(shot)
        whited = whited_pixels(os.path.join(WORK, "doc.png"), shot)
        check(whited > 300, "пиксели данных выжжены", f"закрыто точек: {whited}")

    print("\n== изображение ==")
    r = run(["doc.png"])
    out_img = os.path.join(WORK, "doc_anon.png")
    check(r.returncode == 0, "код возврата 0", r.stderr.strip()[:120])
    check(os.path.exists(out_img), "создан doc_anon.png")
    if os.path.exists(out_img):
        check(Image.open(out_img).size == Image.open(png).size, "размер сохранён")
        whited = whited_pixels(png, out_img)
        check(whited > 300, "пиксели данных выжжены", f"закрыто точек: {whited}")

    print("\n== выбор категорий ==")
    r = run(["doc.md", "--keep", "person", "-o", "keep.md"])
    if os.path.exists(os.path.join(WORK, "keep.md")):
        text = open(os.path.join(WORK, "keep.md"), encoding="utf-8").read()
        check("Иванов Иван Иванович" in text, "--keep person оставляет ФИО")
        check("7707083893" not in text, "остальные категории скрыты")

    r = run(["doc.md", "--ask", "-o", "ask.md"], stdin_text="person\n")
    if os.path.exists(os.path.join(WORK, "ask.md")):
        text = open(os.path.join(WORK, "ask.md"), encoding="utf-8").read()
        check("Иванов Иван Иванович" in text, "--ask учитывает ответ")

    print("\n== отчёт и ошибки ==")
    r = run(["doc.md", "-o", "rep.md", "--report", "report.json"])
    rep_path = os.path.join(WORK, "report.json")
    ok = os.path.exists(rep_path)
    check(ok, "JSON-отчёт создан")
    if ok:
        data = json.load(open(rep_path, encoding="utf-8"))
        cats = data[0].get("categories", {})
        check("person" in cats and "codes" in cats, "в отчёте есть категории", str(list(cats)))

    open(os.path.join(WORK, "bad.xyz"), "w").write("test")
    r = run(["bad.xyz"])
    check(r.returncode != 0, "неподдерживаемый формат отклонён", f"код {r.returncode}")

    r = run(["нет-такого-файла.pdf"])
    check(r.returncode != 0, "отсутствующий файл обработан как ошибка", f"код {r.returncode}")

    failed = [name for ok, name, _ in results if not ok]
    print(f"\nИтог: {len(results) - len(failed)} из {len(results)} проверок пройдено")
    if failed:
        print("Проблемы:")
        for name in failed:
            print("  -", name)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
